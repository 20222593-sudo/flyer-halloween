# Arma export/Agente-Cross-2026.pptx a partir de export/slides/*.png (una imagen por diapositiva)
# y copia el texto de cada <section class="slide"> de presentacion.html en las notas del orador.
import glob, html, re
from pptx import Presentation
from pptx.util import Emu

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
src = open('presentacion.html', encoding='utf-8').read()
sections = re.findall(r'<section class="slide.*?</section>', src, re.S)
for png, sec in zip(sorted(glob.glob('export/slides/slide-*.png')), sections):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(png, 0, 0, prs.slide_width, prs.slide_height)
    text = re.sub(r'<svg.*?</svg>', '', sec, flags=re.S)
    text = html.unescape(re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', text))).strip()
    slide.notes_slide.notes_text_frame.text = text
prs.save('export/Agente-Cross-2026.pptx')
print(len(prs.slides), 'slides')
